# -*- coding: utf-8 -*-
# Part of Odoo. See LICENSE file for full copyright and licensing details.

import base64
import io
from docx import Document
from docx.shared import Pt

from odoo.exceptions import UserError
from odoo.tests.common import TransactionCase, tagged


@tagged('post_install', '-at_install')
class TestContractWordTemplate(TransactionCase):
    """Test suite covering HR Contract Word Template lifecycle and dynamic docx generation."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company_a = cls.env.company
        cls.contract_type = cls.env['hr.contract.type'].create({
            'name': 'Hợp đồng lao động chính thức',
            'sequence': 1,
        })
        cls.employee = cls.env['hr.employee'].create({
            'name': 'Nguyễn Văn A',
            'sex': 'male',
            'work_email': 'nguyenvana@example.com',
            'work_phone': '0901234567',
        })
        # If custom_hr_contract fields exist on employee, populate them
        if hasattr(cls.employee, 'id_card_place'):
            cls.employee.write({
                'id_card_place': 'Cục Cảnh sát QLHC về TTXH',
                'permanent_street': '123 Đường Lê Lợi',
                'permanent_city': 'TP. Hồ Chí Minh',
            })

        # Generate a minimal valid .docx template in memory with Jinja2 placeholders
        doc = Document()
        doc.add_heading('CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM', level=1)
        doc.add_paragraph('HỢP ĐỒNG LAO ĐỘNG SỐ: {{ contract_name }}')
        doc.add_paragraph('Họ và tên người lao động: {{ employee_name }}')
        doc.add_paragraph('Vị trí công việc: {{ job_title }}')
        doc.add_paragraph('Mức lương: {{ wage }} VND (Bằng chữ: {{ wage_in_words }})')
        doc.add_paragraph('Đại diện công ty: {{ company_name }}')

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        cls.template_binary = base64.b64encode(buf.read()).decode('utf-8')

        cls.word_template = cls.env['hr.contract.word.template'].create({
            'name': 'Mẫu HĐLĐ Chính Thức 2026',
            'template_file': cls.template_binary,
            'filename': 'mau_hdld_chinh_thuc.docx',
            'company_id': cls.company_a.id,
        })

        cls.contract = cls.env['hr.contract'].create({
            'name': 'HDLD/2026/0001',
            'employee_id': cls.employee.id,
            'type_id': cls.contract_type.id,
            'wage': 20000000,
            'date_start': '2026-01-01',
            'word_template_id': cls.word_template.id,
            'company_id': cls.company_a.id,
        })

    def test_01_template_creation(self):
        """Verify contract word template fields and default values."""
        self.assertTrue(self.word_template.id)
        self.assertEqual(self.word_template.name, 'Mẫu HĐLĐ Chính Thức 2026')
        self.assertTrue(self.word_template.template_file)
        self.assertTrue(self.word_template.active)

    def test_02_build_docx_context(self):
        """Verify context dictionary mapping employee, contract, and Vietnamese localization."""
        context = self.contract._build_docx_context()
        self.assertEqual(context['employee_name'], 'Nguyễn Văn A')
        self.assertEqual(context['contract_name'], 'HDLD/2026/0001')
        self.assertEqual(context['wage'], '20,000,000')
        self.assertIn('Hai mươi triệu', context['wage_in_words'])
        self.assertEqual(context['company_name'], self.company_a.name)

    def test_03_action_print_docx_contract_success(self):
        """Verify DOCX generation renders cleanly and creates a downloadable attachment."""
        action = self.contract.action_print_docx_contract()
        self.assertEqual(action.get('type'), 'ir.actions.act_url')
        self.assertIn('/web/content/', action.get('url'))
        self.assertIn('download=true', action.get('url'))

        # Check attachment was created and contains valid docx
        attachment = self.env['ir.attachment'].search([
            ('res_model', '=', 'hr.contract'),
            ('res_id', '=', self.contract.id),
        ], order='id desc', limit=1)
        self.assertTrue(attachment)
        self.assertTrue(attachment.name.startswith('HDLD_Nguyen_Van_A_'))
        self.assertTrue(attachment.name.endswith('.docx'))

        # Parse output docx to ensure placeholders were replaced
        out_bytes = io.BytesIO(bytes(attachment.raw))
        out_doc = Document(out_bytes)
        full_text = '\n'.join([p.text for p in out_doc.paragraphs])
        self.assertIn('Họ và tên người lao động: Nguyễn Văn A', full_text)
        self.assertIn('Mức lương: 20,000,000 VND', full_text)
        self.assertIn('Hai mươi triệu', full_text)

    def test_04_action_print_docx_without_template_raises(self):
        """Calling action_print_docx_contract without template raises UserError."""
        contract_no_tpl = self.env['hr.contract'].create({
            'name': 'HDLD/2026/0002',
            'employee_id': self.employee.id,
            'type_id': self.contract_type.id,
            'wage': 15000000,
            'date_start': '2026-02-01',
            'word_template_id': False,
            'company_id': self.company_a.id,
        })
        with self.assertRaises(UserError):
            contract_no_tpl.action_print_docx_contract()

    def test_05_employee_contract_tracking(self):
        """Verify employee contract stat tracking and action_open_contract."""
        self.assertEqual(self.employee.contracts_count, 1)
        self.assertEqual(self.employee.contract_id, self.contract)
        action = self.employee.action_open_contract()
        self.assertEqual(action.get('res_model'), 'hr.contract')
        self.assertEqual(action.get('domain'), [('employee_id', '=', self.employee.id)])
        self.assertEqual(action.get('context', {}).get('default_employee_id'), self.employee.id)
